from pathlib import Path
from typing import Any

from pptx import Presentation


def load_pptx(path: str | Path) -> list[dict[str, Any]]:
    """Extract text slide-by-slide and preserve slide metadata."""
    path = Path(path)
    presentation = Presentation(path)
    slides: list[dict[str, Any]] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = shape.text.strip()
            if text:
                parts.append(text)

        combined = "\n".join(parts).strip()
        if not combined:
            continue

        slides.append({
            "text": combined,
            "metadata": {
                "source_name": path.name,
                "source_type": "pptx",
                "slide": slide_number,
            },
        })

    return slides
